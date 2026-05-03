import sqlite3
from openpyxl import load_workbook

excel_file = 'example_shares.xlsx'
db_file = 'example_shares_excel.db'
counter=0

connection = sqlite3.connect(db_file)
cursor = connection.cursor()

cursor.execute("""
CREATE TABLE IF NOT EXISTS shares (
    title TEXT primary key,
    swx TEXT,
    price REAL,
    dividend REAL,
    volume INTEGER,
    monthly_income TEXT
)
""")

# We want the formula, not just the value, hence data_only=False
workbook = load_workbook(excel_file, data_only=False)
sheet = workbook.active

for row in sheet.iter_rows(min_row=2, values_only=True):
    title, swx, price, dividend, volume, monthly_income = row

    # We don't want the header row
    if title=='Title':
        continue

    counter+=1
    cursor.execute("""
    REPLACE INTO shares (title, swx, price, dividend, volume, monthly_income)
    VALUES (?, ?, ?, ?, ?, ?)
    """, (
        title,
        swx,
        float(price),
        float(dividend),
        int(volume),
        monthly_income
    ))

connection.commit()
connection.close()

print(f'{counter} records have been written into the database')
