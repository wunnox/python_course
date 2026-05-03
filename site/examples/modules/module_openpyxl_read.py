from openpyxl import load_workbook

excel_file = 'example_shares.xlsx'

# Load the workbook
wb = load_workbook(excel_file)
ws = wb.active   # or: wb['Shares'], if named as such

# First row is the header
headers = [cell.value for cell in ws[1]]
print('Header:', headers)

# Read data rows from row 2 onwards
data = []
for row in ws.iter_rows(min_row=2, values_only=True):
    data.append(row)

# Output the read data
for row in data:
    print(row)
